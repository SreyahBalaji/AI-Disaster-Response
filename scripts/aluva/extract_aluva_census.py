
from pathlib import Path
import csv
import openpyxl

source = Path(
    "data/raw/aluva/"
    "2011-IndiaStateDistSbDistTwnWrd-0000.xlsx"
)

output = Path("data/processed/aluva/aluva_census_2011.csv")

wb = openpyxl.load_workbook(
    source, read_only=True, data_only=True
)
ws = wb["Data"]

rows = ws.iter_rows(values_only=True)
headers = next(rows)

# Census geographic identifiers and population columns
columns = [
    "State", "District", "Subdistt",
    "Town/Village", "Ward", "Level",
    "Name", "TRU", "No_HH",
    "TOT_P", "TOT_M", "TOT_F"
]

indices = {name: headers.index(name) for name in columns}

# Identify the Aluva municipality's census town code
aluva_towns = []

for row in rows:
    name = str(row[indices["Name"]] or "").strip()

    if (
        "aluva" in name.lower()
        and row[indices["Level"]] == "TOWN"
    ):
        aluva_towns.append({
            col: row[idx] for col, idx in indices.items()
        })

print("Aluva town matches:")
for town in aluva_towns:
    print(town)

aluva_towns = [
    town for town in aluva_towns
    if str(town["Town/Village"]) == "803286"
    and town["Name"] == "Aluva (M)"
]

if len(aluva_towns) != 1:
    print(
        "STOP: Confirm the correct municipality "
        "before extracting ward records."
    )
else:
    town = aluva_towns[0]
    key_columns = ["State", "District", "Town/Village"]

    wb.close()
    wb = openpyxl.load_workbook(
        source, read_only=True, data_only=True
    )
    ws = wb["Data"]

    extracted = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        same_town = all(
            str(row[indices[col]]) == str(town[col])
            for col in key_columns
        )

        if same_town and row[indices["Level"]] in ("TOWN", "WARD"):
            extracted.append({
                col: row[idx] for col, idx in indices.items()
            })

    output.parent.mkdir(parents=True, exist_ok=True)

    with output.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        writer.writerows(extracted)

    print("\nExtracted records:", len(extracted))
    print("Saved:", output)

    for record in extracted:
        print(
            record["Level"],
            record["Ward"],
            record["Name"],
            record["TOT_P"]
        )

wb.close()
